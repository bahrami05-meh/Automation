# پروزه اتوماسیون بورسی
from pathlib import Path
from docx import Document
from build_jack_project_doc import add_body, add_heading

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
title = '۱۴. توسعه و استفاده روی دو دستگاه'
if not any(p.text == title for p in doc.paragraphs):
    backup = root / '_label_backups' / 'before-two-device.docx'
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    add_heading(doc, title)
    add_body(doc, 'برنامه برای استفادهٔ شخصی و نوبتی روی کامپیوتر محل کار یا لپ‌تاپ منزل آماده می‌شود. کد، مستندات و اسکیل‌ها از طریق GitHub مشترک‌اند؛ محیط پایتون، تنظیمات دستگاه، نشست مرورگر و داده‌های خصوصی مستقل می‌مانند. راهنمای کامل و وضعیت پیاده‌سازی در TWO-DEVICE.md ثبت شده است.')
    add_body(doc, 'مسیر ابزارهای سند نسبت به ریشهٔ پروژه محاسبه می‌شود. setup-local.ps1 محیط محلی را آماده و check-local.py تنظیمات را بررسی می‌کند. میزبان پیشنهادی 127.0.0.1 و پورت پیشنهادی 8765 است. این‌ها مقدمات‌اند و به معنی ساخته‌شدن سرور نیستند. آدرس مخزن و آزمون دستگاه دوم هنوز باقی‌اند.')
    add_body(doc, 'پیش از شروع هر دستگاه نسخهٔ آخر دریافت شود؛ پایان مرحله تغییرات بررسی، ثبت و ارسال شوند. ویرایش هم‌زمان سند Word در دو دستگاه انجام نشود. انتقال خودکار تاریخچهٔ خصوصی و خروجی رمزگذاری‌شده هنوز پیاده‌سازی نشده است.')
    doc.save(path)
assert any(p.text == title for p in Document(path).paragraphs)
print('Two-device document section verified.')
