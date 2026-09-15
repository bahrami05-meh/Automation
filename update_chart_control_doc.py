# پروزه اتوماسیون بورسی
from pathlib import Path
from docx import Document
from build_jack_project_doc import add_body, add_heading

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
title = '۱۷. کنترل نمودار و افزودن اندیکاتور'
if not any(p.text == title for p in doc.paragraphs):
    backup = root / '_label_backups' / 'before-chart-control.docx'
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    add_heading(doc, title)
    add_body(doc, 'ایجنت کنترل نمودار پس از ورود دستی کاربر، نماد، تایم‌فریم و اندیکاتور درخواستی را در سایت نموداری تأییدشده باز و تنظیم می‌کند. نام سایت، زمان دریافت، نماد، تایم‌فریم، نام اندیکاتور و پارامترهای اعمال‌شده در گزارش ثبت می‌شوند.')
    add_body(doc, 'انتخاب خودکار اندیکاتور فقط از فهرست مجاز و با دلیل ثبت‌شده بر پایهٔ نوع تحلیل، تایم‌فریم و دادهٔ موجود انجام می‌شود. RSI، MACD، میانگین متحرک، ایچیموکو و ATR نمونه‌های اولیه‌اند. اگر دادهٔ OHLCV معتبر وجود دارد، محاسبهٔ مستقل در هستهٔ پایتون مرجع کنترل مقدارهای نمودار است.')
    add_body(doc, 'اگر منوی سایت تغییر کند، اندیکاتور در دسترس نباشد، نمودار خوانا نباشد یا دادهٔ نمودار با دادهٔ OHLCV ناسازگار باشد، ایجنت باید DATA_BLOCKED یا اختلاف منبع اعلام کند. ورود، CAPTCHA و احراز هویت دومرحله‌ای با کاربر است. تعامل ایجنت با صفحهٔ کارگزاری فقط‌خواندنی است و هیچ اقدام مالی انجام نمی‌شود.')
    doc.save(path)
assert any(p.text == title for p in Document(path).paragraphs)
print('Chart-control report section verified.')
