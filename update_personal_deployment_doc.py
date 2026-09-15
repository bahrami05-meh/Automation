# پروزه اتوماسیون بورسی
from pathlib import Path
from docx import Document
from build_jack_project_doc import add_body, add_heading

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
replacements = {
    'نام و نوع سامانهٔ کارگزاری و این‌که آیا مسیر رسمی API فقط‌خواندنی در اختیار دارد یا خیر.': 'نام و آدرس سامانهٔ بورسی برای بررسی دریافت از مرورگر؛ نبود API توسط مالک پروژه اعلام شده است.',
    'محل استقرار: کامپیوتر شخصی، سرور خصوصی یا محیط ابری؛ با توجه به دادهٔ مالی حساس.': 'انتخاب مدل هوش مصنوعی و اجرای محلی یا استفاده از API مدل؛ محل اجرای برنامه، کامپیوتر محل کار یا لپ‌تاپ منزل تعیین شده است.',
}
backup = root / '_label_backups' / 'before-personal-deployment.docx'
if not backup.exists():
    backup.write_bytes(path.read_bytes())
for paragraph in doc.paragraphs:
    if paragraph.text in replacements:
        value = replacements[paragraph.text]
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ''
title = 'تصمیم‌های قطعی دریافت و اجرا'
if not any(p.text == title for p in doc.paragraphs):
    anchor = next(p for p in doc.paragraphs if p.text == '۵. معماری منطقی')
    fragment = Document()
    add_heading(fragment, title, 2)
    text = (root / 'PROJECT.md').read_text(encoding='utf-8').split('## ' + title + '\n', 1)[1].split('\n## ', 1)[0]
    for body in text.strip().split('\n\n'):
        add_body(fragment, body)
    previous = anchor._p
    for p in fragment.paragraphs:
        previous.addnext(p._p)
        previous = p._p
doc.save(path)
check = Document(path)
assert check.paragraphs[0].text == 'پروزه اتوماسیون بورسی'
assert any(p.text == title for p in check.paragraphs)
assert not any(p.text in replacements for p in check.paragraphs)
print('Deployment decisions updated and verified.')
