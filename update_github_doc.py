# پروزه اتوماسیون بورسی
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
replacements = {
    'نشانی مخزن هنوز ارائه نشده و ارسال به GitHub انجام نشده است.': 'مخزن معرفی‌شدهٔ مالک پروژه https://github.com/bahrami05-meh/Automation.git و شاخهٔ پایه main است. وضعیت ارسال از تاریخچهٔ Git بررسی شود.',
    'آدرس مخزن و آزمون دستگاه دوم هنوز باقی‌اند.': 'مخزن GitHub تعیین شده است؛ آزمون دستگاه دوم هنوز باقی است.',
}
changed = False
for p in doc.paragraphs:
    updated = p.text
    for before, after in replacements.items():
        updated = updated.replace(before, after)
    if updated != p.text:
        p.runs[0].text = updated
        for run in p.runs[1:]:
            run.text = ''
        changed = True
if changed:
    backup = root / '_label_backups' / 'before-github.docx'
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    doc.save(path)
assert any('github.com/bahrami05-meh/Automation.git' in p.text for p in Document(path).paragraphs)
print('GitHub repository recorded in project document.')
