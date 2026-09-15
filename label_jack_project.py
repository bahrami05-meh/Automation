from pathlib import Path
import re
import sys
from docx import Document

LABEL = 'پروزه اتوماسیون بورسی'
ROOT = Path(__file__).resolve().parent
SKILLS = ROOT / 'skills'
NAMES = ['security-best-practices', 'security-threat-model', 'jupyter-notebook', 'iran-market-analysis', 'Made GL-5']

if '--skills' in sys.argv:
    for name in NAMES:
        path = SKILLS / name / 'SKILL.md'
        original = path.read_bytes()
        text = original.decode('utf-8-sig')
        match = re.match(r'\A---\r?\n.*?\r?\n---(?:\r?\n|$)', text, re.S)
        if not match:
            raise ValueError(f'Missing frontmatter: {path}')
        start, body = text[:match.end()], text[match.end():]
        if body.lstrip().startswith(LABEL):
            continue
        newline = '\r\n' if '\r\n' in text else '\n'
        updated = start + newline + LABEL + newline + body
        backup = ROOT / '_label_backups' / name / 'SKILL.md'
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            backup.write_bytes(original)
        path.write_bytes(updated.encode('utf-8'))
        assert path.read_text(encoding='utf-8').split('---', 2)[2].strip().startswith(LABEL)
        print(name + ': labeled; frontmatter preserved')
else:
    sys.path.insert(0, str(ROOT / 'document_tools'))
    from fa_docx import set_rtl_paragraph
    path = ROOT / 'مستند-پروژه-جک.docx'
    doc = Document(path)
    if doc.paragraphs[0].text != LABEL:
        backup = ROOT / '_label_backups' / path.name
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            backup.write_bytes(path.read_bytes())
        paragraph = doc.paragraphs[0].insert_paragraph_before(LABEL)
        set_rtl_paragraph(paragraph)
        doc.save(path)
    assert Document(path).paragraphs[0].text == LABEL
    print('Document first paragraph verified')
