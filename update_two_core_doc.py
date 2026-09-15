from pathlib import Path
from docx import Document
from build_jack_project_doc import add_body, add_heading

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
source = (root / 'build_jack_project_doc.py').read_text(encoding='utf-8')
start = source.index('    add_heading(doc, "معماری دو هسته‌ای", 2)')
end = source.index('    add_table(doc, ["لایه"', start)
section_code = source[start:end]
if not any(p.text == 'معماری دو هسته‌ای' for p in doc.paragraphs):
    anchor = next(p for p in doc.paragraphs if p.text == '۵. معماری منطقی')
    temp = Document()
    exec(compile(section_code.replace('    ', '', 1) if False else '\n'.join(line[4:] for line in section_code.splitlines()), '<architecture-section>', 'exec'), {'doc': temp, 'add_body': add_body, 'add_heading': add_heading})
    backup = root / '_label_backups' / 'before-two-core.docx'
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    previous = anchor._p
    for paragraph in temp.paragraphs:
        previous.addnext(paragraph._p)
        previous = paragraph._p
    doc.save(path)
check = Document(path)
assert check.paragraphs[0].text == 'پروزه اتوماسیون بورسی'
assert any(p.text == 'معماری دو هسته‌ای' for p in check.paragraphs)
print('Two-core architecture added; document first line preserved.')
