"""Persian/RTL helpers for python-docx documents.

Word and LibreOffice perform Unicode shaping and bidi themselves, so unlike
PDF generation no reshaping is needed. What DOCX needs instead is the correct
OOXML properties — and missing any one of them is the classic way Persian
documents come out left-aligned, wrongly ordered, or in a fallback font:

- paragraph direction:  <w:bidi/> in pPr           -> set_rtl_paragraph()
- run direction:        <w:rtl/> in rPr            -> handled by the same call
- complex-script font:  <w:rFonts w:cs="..."/>     -> set_font() (w:ascii
  alone does NOT apply to Persian text; w:szCs sizes it)
- table column order:   <w:bidiVisual/> in tblPr   -> set_rtl_table()
  (columns then display right-to-left; keep writing them in logical order)

Usage with the plugin venv (gap_runtime.py py):

    from docx import Document
    from fa_docx import add_rtl_paragraph, set_rtl_table, persian_digits
    doc = Document()
    add_rtl_paragraph(doc, "گزارش فروش", style="Title")
    table = doc.add_table(rows=1, cols=3)
    set_rtl_table(table)
"""

from docx.oxml.ns import qn
from docx.shared import Pt

PERSIAN_DIGITS = str.maketrans("0123456789,", "۰۱۲۳۴۵۶۷۸۹٬")

DEFAULT_PERSIAN_FONT = "Vazirmatn"
# Fallbacks present on most systems if Vazirmatn is not installed for the
# renderer: macOS ships Geeza Pro; Windows ships Tahoma with Arabic script.
PERSIAN_FALLBACK_FONTS = ("Vazirmatn", "Geeza Pro", "Tahoma", "Arial")


def persian_digits(value):
    """Localize ASCII digits and thousands separators: 1,234 -> ۱٬۲۳۴."""
    return str(value).translate(PERSIAN_DIGITS)


def _get_or_add(parent, tag):
    child = parent.find(qn(tag))
    if child is None:
        child = parent.makeelement(qn(tag), {})
        parent.append(child)
    return child


def set_font(run, name=DEFAULT_PERSIAN_FONT, size=None, cs_only=False):
    """Set a run's font incl. the complex-script (w:cs) slot Persian uses."""
    rPr = run._element.get_or_add_rPr()
    fonts = _get_or_add(rPr, "w:rFonts")
    if not cs_only:
        fonts.set(qn("w:ascii"), name)
        fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
        szCs = _get_or_add(rPr, "w:szCs")
        szCs.set(qn("w:val"), str(int(size * 2)))


def set_rtl_paragraph(paragraph, font=DEFAULT_PERSIAN_FONT, size=None):
    """Make one paragraph right-to-left (direction + run rtl + cs font)."""
    pPr = paragraph._p.get_or_add_pPr()
    _get_or_add(pPr, "w:bidi")
    for run in paragraph.runs:
        rPr = run._element.get_or_add_rPr()
        _get_or_add(rPr, "w:rtl")
        if font:
            set_font(run, font, size)
    return paragraph


def add_rtl_paragraph(container, text, style=None, font=DEFAULT_PERSIAN_FONT,
                      size=None):
    """add_paragraph + set_rtl_paragraph in one call."""
    paragraph = container.add_paragraph(text, style=style)
    return set_rtl_paragraph(paragraph, font=font, size=size)


def set_rtl_table(table, font=DEFAULT_PERSIAN_FONT):
    """Display table columns right-to-left and make cell text RTL.

    Keep building rows in logical order (first column = rightmost).
    """
    tblPr = table._tbl.tblPr
    _get_or_add(tblPr, "w:bidiVisual")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_rtl_paragraph(paragraph, font=font)
    return table


def set_rtl_styles(document, font=DEFAULT_PERSIAN_FONT):
    """Point the document's base styles at a Persian-capable font.

    Word styles carry their own rFonts; setting Normal + headings once keeps
    every non-directly-formatted run consistent.
    """
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2",
                       "Heading 3", "Heading 4"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        rPr = style.element.get_or_add_rPr()
        fonts = _get_or_add(rPr, "w:rFonts")
        fonts.set(qn("w:ascii"), font)
        fonts.set(qn("w:hAnsi"), font)
        fonts.set(qn("w:cs"), font)
    return document
