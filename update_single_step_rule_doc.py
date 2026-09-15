# پروزه اتوماسیون بورسی
from pathlib import Path
from docx import Document
from build_jack_project_doc import add_body, add_heading

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
title = '۱۶. نتیجهٔ نهایی و پلهٔ محدود'
if not any(p.text == title for p in doc.paragraphs):
    backup = root / '_label_backups' / 'before-single-step-rule.docx'
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    add_heading(doc, title)
    add_body(doc, 'نظر نهایی جک برای هر نماد برآیند شواهد ایجنت‌های تابلو، تکنیکال، بنیادی، رویدادها و ریسک است؛ جمع سادهٔ رأی‌ها نیست. خروجی کاربرپسند برای هر نماد یکی از چهار حالت خرید آزمایشی، فروش آزمایشی، نگهداری / واچ‌لیست یا عدم اعلام نظر است. حالت آخر در دادهٔ ناقص، قدیمی، متناقض یا ناکافی اجباری است.')
    add_body(doc, 'خرید یا فروش فقط یک پلهٔ محدود دارد. جک حداکثر اندازهٔ همان پله را بر اساس بودجهٔ ریسک، فاصله تا حد زیان، کارمزد، لغزش، سرمایهٔ آزاد، محدودیت تمرکز و نقدشوندگی محاسبه و نمایش می‌دهد. افزایش خودکار موقعیت، سفارش چندپله‌ای، میانگین کم‌کردن و اجرای مالی واقعی خارج از دامنه است.')
    add_body(doc, 'نتیجه شامل شرط تأیید، شرط ابطال، شواهد موافق و مخالف، منبع و زمان داده و تأیید انسانی است. این نتیجه فقط برای محیط آزمایشی است و به سفارش واقعی تبدیل نمی‌شود.')
    doc.save(path)
assert any(p.text == title for p in Document(path).paragraphs)
print('Single-step rule verified.')
