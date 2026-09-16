# پروزه اتوماسیون بورسی
from pathlib import Path
from docx import Document
from build_jack_project_doc import add_body, add_heading

root = Path(__file__).resolve().parent
path = root / 'مستند-پروژه-جک.docx'
doc = Document(path)
title = 'تقویم و زمان‌بندی بورس ایران'
if not any(p.text == title for p in doc.paragraphs):
    backup = root / '_label_backups' / 'before-market-schedule.docx'
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    add_heading(doc, title, 2)
    add_body(doc, 'برای بازار سهام بورس و فرابورس ایران، زمان‌بندی پایهٔ جک شنبه تا چهارشنبه، جلسهٔ اصلی معاملات از ۰۹:۰۰ تا ۱۲:۳۰ به منطقهٔ زمانی Asia/Tehran است. پیش‌گشایش حدود ۰۸:۴۵ و مرحلهٔ معاملات پایانی برای برخی نمادها ۱۲:۴۵ تا ۱۳:۰۰ است. این زمان‌ها باید در تنظیمات قابل‌اصلاح باشند؛ تعطیلات رسمی، توقف نماد، اطلاعیهٔ بازار و تغییر مقررات بر برنامهٔ پایه مقدم‌اند.')
    add_body(doc, 'زمان‌بند فقط در پنجره‌های تعیین‌شده تحلیل تازه و پایش بازار را اجرا می‌کند. بیرون از ساعت بازار، آخرین دادهٔ بازار زنده تلقی نمی‌شود؛ جک فقط گزارش‌های قبلی را نمایش می‌دهد یا برای اجرای بعدی آماده می‌شود. در روز تعطیل یا وقتی تقویم/وضعیت بازار نامشخص است، اجرای تحلیلی با MARKET_CLOSED یا DATA_BLOCKED ثبت می‌شود.')
    doc.save(path)
assert any(p.text == title for p in Document(path).paragraphs)
print('Market schedule section verified.')
